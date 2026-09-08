import pandas as pd
from openpyxl import Workbook, load_workbook

from app.core.active_profile_store import ActiveProfileStore
from app.core.output_profile import OutputProfileStore
from app.core.processing_engine import ProcessingEngine
from app.core import manim_output_service


def _netsis_toplu_xlsx_template(project_root):
    """Netsis'in sağladığı boş XLSX şablonunun test karşılığını kurar."""
    profile = OutputProfileStore(project_root / "config").get("netsis_toplu")
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Sheet1"
    worksheet.append(profile.headers())
    workbook.save(project_root / "templates" / profile.template_file)


def test_toplu_netsis_sablonu_banka_kodlariyla_bolge_basina_tek_dosya_yazar(
    synthetic_project, monkeypatch
):
    manim_path, tahsilat_path, customer_path, project_root = synthetic_project
    rows = pd.read_excel(manim_path)
    rows.loc[0, "Banka"] = "Ziraat"
    rows.loc[0, "Tutar"] = 1200.0
    rows.loc[0, "Açıklama"] = "ZIRAAT HAVALE"
    rows.loc[1, "Banka"] = "Garanti"
    rows.loc[1, "Dekont Durumu"] = "Aktarıldı"
    rows.loc[1, "Karşı Hesap Adı"] = "ABC LTD"
    rows.loc[1, "Karşı Hesap Kodu"] = "ABC001"
    rows.loc[1, "Tutar"] = 1300.0
    rows.loc[1, "Açıklama"] = "GARANTI HAVALE"
    rows = rows.iloc[:2]
    rows.to_excel(manim_path, index=False)

    _netsis_toplu_xlsx_template(project_root)
    class TestWriter:
        def __init__(self, profile):
            self.profile = profile

        def write(self, records, output_path):
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Sheet1"
            sheet.append(self.profile.headers())
            for record in records:
                sheet.append([
                    getattr(record, column.field) if column.source_kind == "field"
                    else column.value
                    for column in self.profile.columns
                ])
            destination = output_path.with_suffix(self.profile.output_extension)
            workbook.save(destination)
            return destination

        def close(self):
            pass

    monkeypatch.setattr(manim_output_service, "NetsisWriter", TestWriter)
    ActiveProfileStore(project_root).set_output_profile_id("netsis_toplu")
    result = ProcessingEngine([manim_path, tahsilat_path, customer_path], project_root).run()

    netsis_files = [path for path in result.created_files if "BODRUM" in path.name]
    assert len(netsis_files) == 1
    assert "GARANTI" not in netsis_files[0].name
    assert "ZIRAAT" not in netsis_files[0].name

    assert netsis_files[0].suffix == ".xlsx"
    workbook = load_workbook(netsis_files[0], data_only=True)
    sheet = workbook["Sheet1"]
    assert sheet.cell(1, 1).value == "Banka Hes.Kodu(*)"
    assert sheet.cell(1, 19).value == "Muh.Ref.Kod(*)"
    rows = list(sheet.iter_rows(min_row=2, values_only=True))
    assert {row[0] for row in rows} == {
        "BANK-G-01", "BANK-Z-01"
    }
    assert {row[1] for row in rows} == {0}
    assert {row[18] for row in rows} == {"G01"}
    assert {row[19] for row in rows} == {0}
    assert {row[20] for row in rows} == {"HV"}


def test_toplu_sablonda_bm_kodu_tanimli_olmayan_banka_bos_kodla_yazilmaz(
    synthetic_project,
):
    manim_path, tahsilat_path, customer_path, project_root = synthetic_project
    rows = pd.read_excel(manim_path).iloc[:1].copy()
    rows.loc[0, "Banka"] = "Akbank"
    rows.loc[0, "Dekont Durumu"] = "Aktarıldı"
    rows.loc[0, "Karşı Hesap Kodu"] = "ABC001"
    rows.to_excel(manim_path, index=False)

    ActiveProfileStore(project_root).set_output_profile_id("netsis_toplu")
    result = ProcessingEngine([manim_path, tahsilat_path, customer_path], project_root).run()

    assert result.produced_netsis_records == 0
    assert result.unresolved == 1
    assert result.review_file is not None
    review = pd.read_excel(result.review_file)
    assert "BM banka hesap kodu tanımlı değil" in review.loc[0, "Neden"]
